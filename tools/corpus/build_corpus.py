#!/usr/bin/env python3
"""构建规模测试语料：真实书目打底，扩展到 5k / 10k / 50k。

数据来源（两份公开数据集，均从 ModelScope 下载，2026-10-29 实测可达）：
  - YU0606777/DOUBAN-BOOKS-TOP250         书名/作者/出版社/年份/评分
  - G13565855469/dangdang_24h_book_dataset 书名/作者/出版社/年份/价格

真实字段只有：书名、作者、出版社、出版年。
生成字段：ISBN（占位值，源数据不含）、馆藏位置、在馆状态、简介（按模板生成）。
扩展部分每条都能追溯到一本真实书目，只在书名后加版次/卷次后缀。

用法：
    python tools/corpus/build_corpus.py --dangdang 当当图书.xlsx --douban TOP250.xlsx \
        --out results/dataset
"""
import argparse
import hashlib
import json
import os
import random

from openpyxl import load_workbook

SEED = 20261029

# 书名关键词 -> 领域，用于生成简介，让不同领域的书在向量空间里能分开
DOMAINS = [
    (("Java", "Python", "Go", "Rust", "C++", "JavaScript", "Node", "编程", "算法",
      "代码", "数据库", "Redis", "MySQL", "Linux", "Kubernetes", "Docker", "架构",
      "重构", "编译", "网络", "操作系统", "深度学习", "机器学习", "大模型"), "计算机技术"),
    (("数学", "几何", "代数", "物理", "化学", "生物", "科学", "统计"), "自然科学"),
    (("历史", "文明", "帝国", "战争", "王朝", "简史"), "历史"),
    (("经济", "金融", "投资", "货币", "管理", "商业", "营销"), "经济管理"),
    (("小说", "文集", "诗", "散文", "记", "传", "童话", "侦探", "悬疑"), "文学"),
    (("心理", "思维", "认知", "情绪", "社会"), "心理与社会"),
    (("教育", "学习", "考试", "解题", "辅导"), "教育"),
    (("健康", "养生", "饮食", "菜", "健身"), "生活"),
]

TEMPLATES = [
    "《{title}》由{author}著，{publisher}出版。本书围绕{domain}领域展开，"
    "系统梳理核心概念与实践方法，适合希望建立完整知识框架的读者。",
    "《{title}》是{domain}方向的入门与进阶读物，作者{author}结合多年经验，"
    "从基础原理讲到实际应用，配有大量案例，由{publisher}出版。",
    "{author}所著《{title}》，{publisher}出版。全书聚焦{domain}的关键问题，"
    "论述严谨、结构清晰，被多所高校列为参考书目。",
    "《{title}》（{author}著，{publisher}）以{domain}为主题，"
    "既讲理论推导也讲落地方法，适合作为课程教材或自学读本。",
]

VARIANTS = ["第{n}版", "修订版", "典藏版", "插图版", "第{n}卷", "新版", "精装版"]


def load_dangdang(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}
    out = []
    for row in rows[1:]:
        title = (row[idx["书名"]] or "").strip() if len(row) > idx["书名"] else ""
        if not title:
            continue
        author = (row[idx["作者"]] or "").strip()
        author = author.replace("作者:", "").split("著")[0].strip() or "佚名"
        publisher = (row[idx["出版社"]] or "").strip() or "未知出版社"
        year = (row[idx["出版时间"]] or "").strip().replace("出版时间:", "")
        out.append({"title": title, "author": author, "publisher": publisher,
                    "year": year, "source": "dangdang-24h-bestsellers"})
    return out


def load_douban(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}
    out = []
    for row in rows[1:]:
        title = (row[idx["NAME"]] or "").strip() if len(row) > idx["NAME"] else ""
        if not title:
            continue
        raw = (row[idx["ZUOZHE"]] or "").strip()
        parts = [p.strip() for p in raw.split("/")]
        author = parts[0] if parts and parts[0] else "佚名"
        publisher = parts[1] if len(parts) > 1 else "未知出版社"
        year = parts[2] if len(parts) > 2 else ""
        out.append({"title": title, "author": author, "publisher": publisher,
                    "year": year, "source": "douban-top250"})
    return out


def merge(dangdang, douban):
    """按书名去重，豆瓣优先（作者/出版社字段更规整）。"""
    seen, merged = set(), []
    for rec in douban + dangdang:
        key = rec["title"].replace(" ", "")
        if key in seen:
            continue
        seen.add(key)
        merged.append(rec)
    return merged


def domain_of(title):
    for keywords, domain in DOMAINS:
        if any(k in title for k in keywords):
            return domain
    return "人文社科"


def make_isbn(seed_text):
    digest = hashlib.sha1(seed_text.encode("utf-8")).hexdigest()
    digits = "".join(c for c in digest if c.isdigit())[:10].ljust(10, "0")
    return "978" + digits


def make_description(rec, rng):
    return rng.choice(TEMPLATES).format(
        title=rec["title"], author=rec["author"],
        publisher=rec["publisher"], domain=domain_of(rec["title"]))


def make_book(rec, index, rng):
    title = rec["title"]
    if index >= 0:
        variant = rng.choice(VARIANTS).format(n=(index // 997) % 9 + 2)
        title = f"{title}（{variant}）"
    isbn = make_isbn(f"{rec['title']}#{index}")
    return {
        "isbn": isbn,
        "title": title[:200],
        "author": rec["author"][:100],
        # 与 BookStatus 枚举名一致，才能直接被 DataInitializer 读入
        "status": "AVAILABLE" if rng.random() < 0.8 else "BORROWED",
        "location": f"{rng.choice('ABCDE')}区-{rng.randint(1, 9)}排-{rng.randint(1, 40):02d}号",
        "description": make_description(rec, rng),
    }


def build(real, total, rng):
    books, seen_isbn = [], set()
    for i in range(total):
        rec = real[i % len(real)]
        book = make_book(rec, -1 if i < len(real) else i, rng)
        while book["isbn"] in seen_isbn:  # ISBN 冲突时加盐重算
            book["isbn"] = make_isbn(book["title"] + str(rng.random()))
        seen_isbn.add(book["isbn"])
        books.append(book)
    return books


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dangdang", help="当当畅销书 xlsx；与 --douban 一起提供时重新解析原始数据")
    ap.add_argument("--douban", help="豆瓣 TOP250 xlsx")
    ap.add_argument("--from-catalog", help="直接读已生成的 catalog-real.json，跳过原始解析")
    ap.add_argument("--out", default="results/dataset")
    ap.add_argument("--sizes", default="5000,10000,50000")
    args = ap.parse_args()

    os.makedirs(args.out, exist_ok=True)
    if args.from_catalog:
        with open(args.from_catalog, encoding="utf-8") as f:
            real = json.load(f)
    elif args.dangdang and args.douban:
        real = merge(load_dangdang(args.dangdang), load_douban(args.douban))
    else:
        ap.error("要么给 --from-catalog，要么同时给 --dangdang 和 --douban")
    print(f"真实书目去重后：{len(real)} 条")

    with open(os.path.join(args.out, "catalog-real.json"), "w", encoding="utf-8") as f:
        json.dump(real, f, ensure_ascii=False, indent=1)

    manifest = {
        "generated_by": "tools/corpus/build_corpus.py",
        "seed": SEED,
        "real_book_count": len(real),
        "sources": [
            {"name": "DOUBAN-BOOKS-TOP250", "platform": "ModelScope",
             "namespace": "YU0606777", "license_note": "公开数据集，仅使用书名/作者/出版社/年份等事实字段"},
            {"name": "dangdang_24h_book_dataset", "platform": "ModelScope",
             "namespace": "G13565855469", "license_note": "同上"},
        ],
        "real_fields": ["title", "author", "publisher", "year"],
        "generated_fields": ["isbn", "status", "location", "description"],
        "expansion_rule": "第 i 条取真实书目 i % N，书名追加版次/卷次后缀，"
                          "简介按领域模板生成；每条均可追溯到一本真实书目",
        "scales": {},
    }

    for size in [int(s) for s in args.sizes.split(",")]:
        rng = random.Random(SEED + size)
        books = build(real, size, rng)
        path = os.path.join(args.out, f"books-{size}.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(books, f, ensure_ascii=False)
        manifest["scales"][str(size)] = {
            "file": f"books-{size}.json",
            "count": len(books),
            "real_prefix_count": min(size, len(real)),
            "synthetic_count": max(0, size - len(real)),
        }
        print(f"生成 {path}: {len(books)} 条（真实前缀 {min(size, len(real))} 条）")

    with open(os.path.join(args.out, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
